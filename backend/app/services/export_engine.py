import os
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment


# ============================================================
# CONSTANTS
# ============================================================

MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 10

HEADER_FONT = Font(
    bold=True,
    size=11
)

TITLE_FONT = Font(
    bold=True,
    size=16
)

SUBTITLE_FONT = Font(
    bold=True,
    size=11
)

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


# ============================================================
# GENERAL HELPERS
# ============================================================

def clean_value(value):
    """
    Clean a value before writing it to Excel.
    """

    if value is None:
        return ""

    if isinstance(value, str):
        value = value.replace("\x00", "")
        value = value.replace("\r\n", "\n")
        value = value.strip()

        # Normalize excessive spaces
        value = re.sub(
            r"[ \t]+",
            " ",
            value
        )

        return value

    return value


def normalize_rows(rows):
    """
    Convert incoming rows into a clean list of lists.
    """

    if not rows:
        return []

    normalized = []

    max_columns = 0

    for row in rows:

        if row is None:
            continue

        if not isinstance(row, (list, tuple)):
            row = [row]

        row = [
            clean_value(value)
            for value in row
        ]

        max_columns = max(
            max_columns,
            len(row)
        )

        normalized.append(row)

    # Make every row the same length
    for index in range(
        len(normalized)
    ):

        row = normalized[index]

        if len(row) < max_columns:

            normalized[index] = row + (
                [""] *
                (max_columns - len(row))
            )

    return normalized


def safe_sheet_name(name):
    """
    Excel sheet names cannot contain:
    : \ / ? * [ ]
    """

    if not name:
        name = "Data"

    name = str(name)

    name = re.sub(
        r"[:\\/?*\[\]]",
        "_",
        name
    )

    name = name.strip()

    if not name:
        name = "Data"

    return name[:31]


def make_unique_headers(headers):
    """
    Make duplicate column names unique.

    Example:

    Date
    Date
    Amount

    becomes:

    Date
    Date_2
    Amount
    """

    result = []
    used = {}

    for index, header in enumerate(headers):

        header = clean_value(header)

        if not header:

            header = (
                f"Column_{index + 1}"
            )

        base = str(header)

        if base not in used:

            used[base] = 1
            result.append(base)

        else:

            used[base] += 1

            result.append(
                f"{base}_{used[base]}"
            )

    return result


def convert_possible_number(value):
    """
    Convert obvious numeric values to actual
    Excel numbers.

    Does NOT aggressively convert arbitrary text.
    """

    if value is None:
        return None

    if isinstance(
        value,
        (int, float)
    ):
        return value

    if not isinstance(
        value,
        str
    ):
        return value

    text = value.strip()

    if not text:
        return ""

    # Remove common currency symbols
    cleaned = text.replace(
        "₹",
        ""
    ).replace(
        "$",
        ""
    ).replace(
        "€",
        ""
    ).replace(
        "£",
        ""
    )

    cleaned = cleaned.strip()

    # Negative number in parentheses
    negative = False

    if (
        cleaned.startswith("(")
        and cleaned.endswith(")")
    ):

        negative = True

        cleaned = cleaned[1:-1]

    # Remove thousands separators
    numeric = cleaned.replace(
        ",",
        ""
    )

    # Don't convert IDs / reference numbers
    # that contain non-numeric characters.
    if re.fullmatch(
        r"-?\d+(\.\d+)?",
        numeric
    ):

        try:

            number = float(
                numeric
            )

            if negative:
                number = -number

            if number.is_integer():

                return int(number)

            return number

        except ValueError:
            pass

    return value


def detect_date(value):
    """
    Detect common date formats.
    """

    if not isinstance(
        value,
        str
    ):
        return None

    value = value.strip()

    formats = [
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%d-%m-%y",
        "%d/%m/%y",
        "%d-%b-%Y",
        "%d-%b-%y",
        "%d-%B-%Y",
        "%d-%B-%y",
        "%Y-%m-%d",
    ]

    for date_format in formats:

        try:

            return datetime.strptime(
                value,
                date_format
            )

        except ValueError:
            continue

    return None


# ============================================================
# HEADER / DATA PREPARATION
# ============================================================

def prepare_table(
    rows,
    headers=None
):
    """
    Prepare rows and headers for Excel.

    If headers are not supplied, the first row
    is treated as the header.
    """

    rows = normalize_rows(
        rows
    )

    if not rows:
        return [], []

    if headers is not None:

        headers = make_unique_headers(
            headers
        )

        column_count = len(
            headers
        )

        data_rows = []

        for row in rows:

            if len(row) < column_count:

                row = row + (
                    [""] *
                    (column_count - len(row))
                )

            elif len(row) > column_count:

                row = row[:column_count]

            data_rows.append(row)

        return headers, data_rows

    # First row becomes header
    headers = make_unique_headers(
        rows[0]
    )

    data_rows = rows[1:]

    return headers, data_rows


# ============================================================
# EXCEL FORMATTING
# ============================================================

def apply_header_style(
    worksheet,
    header_row=1
):
    """
    Style Excel header.
    """

    for cell in worksheet[
        header_row
    ]:

        cell.font = HEADER_FONT

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = THIN_BORDER


def apply_data_style(
    worksheet,
    start_row=2
):
    """
    Style all data cells.
    """

    for row in worksheet.iter_rows(
        min_row=start_row
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            cell.border = THIN_BORDER


def detect_column_type(
    values
):
    """
    Determine whether a column contains
    dates, numbers or text.
    """

    non_empty = [
        value
        for value in values
        if value not in (
            "",
            None
        )
    ]

    if not non_empty:
        return "text"

    date_count = 0
    number_count = 0

    for value in non_empty:

        if isinstance(
            value,
            datetime
        ):
            date_count += 1
            continue

        if isinstance(
            value,
            (int, float)
        ):
            number_count += 1
            continue

        if detect_date(
            str(value)
        ):
            date_count += 1
            continue

        converted = convert_possible_number(
            value
        )

        if isinstance(
            converted,
            (int, float)
        ):
            number_count += 1

    total = len(
        non_empty
    )

    if date_count / total >= 0.7:
        return "date"

    if number_count / total >= 0.7:
        return "number"

    return "text"


def format_data_types(
    worksheet,
    start_row=2
):
    """
    Convert obvious dates/numbers to Excel-native
    data types.
    """

    max_column = worksheet.max_column
    max_row = worksheet.max_row

    for column in range(
        1,
        max_column + 1
    ):

        values = []

        for row in range(
            start_row,
            max_row + 1
        ):

            values.append(
                worksheet.cell(
                    row=row,
                    column=column
                ).value
            )

        column_type = detect_column_type(
            values
        )

        for row in range(
            start_row,
            max_row + 1
        ):

            cell = worksheet.cell(
                row=row,
                column=column
            )

            value = cell.value

            if column_type == "date":

                parsed = detect_date(
                    str(value)
                )

                if parsed:

                    cell.value = parsed

                    cell.number_format = (
                        "dd-mm-yyyy"
                    )

            elif column_type == "number":

                converted = convert_possible_number(
                    value
                )

                if isinstance(
                    converted,
                    (int, float)
                ):

                    cell.value = converted

                    cell.number_format = (
                        '#,##0.00'
                    )


def auto_size_columns(
    worksheet
):
    """
    Automatically size columns without allowing
    extremely large widths.
    """

    for column_cells in worksheet.columns:

        max_length = 0

        column_index = (
            column_cells[0].column
        )

        column_letter = get_column_letter(
            column_index
        )

        for cell in column_cells:

            value = cell.value

            if value is None:
                continue

            # Don't calculate enormous width
            # from very long descriptions.
            text = str(value)

            longest_line = max(
                (
                    len(line)
                    for line in text.split("\n")
                ),
                default=0
            )

            max_length = max(
                max_length,
                longest_line
            )

        width = max(
            MIN_COLUMN_WIDTH,
            min(
                max_length + 2,
                MAX_COLUMN_WIDTH
            )
        )

        worksheet.column_dimensions[
            column_letter
        ].width = width


def apply_row_heights(
    worksheet,
    start_row=2
):
    """
    Give rows enough height for wrapped text.
    """

    for row in range(
        start_row,
        worksheet.max_row + 1
    ):

        max_lines = 1

        for column in range(
            1,
            worksheet.max_column + 1
        ):

            value = worksheet.cell(
                row=row,
                column=column
            ).value

            if value is None:
                continue

            text = str(value)

            line_count = max(
                text.count("\n") + 1,
                (len(text) // 70) + 1
            )

            max_lines = max(
                max_lines,
                min(line_count, 6)
            )

        worksheet.row_dimensions[
            row
        ].height = min(
            max_lines * 15,
            90
        )


# ============================================================
# EXCEL TABLE
# ============================================================

def add_excel_table(
    worksheet,
    table_name="ConvertedData"
):
    """
    Add a real Excel table with filters.
    """

    if worksheet.max_row < 2:
        return

    if worksheet.max_column < 1:
        return

    # Excel table names must contain letters,
    # numbers and underscores.
    table_name = re.sub(
        r"[^A-Za-z0-9_]",
        "_",
        table_name
    )

    if not table_name:
        table_name = "ConvertedData"

    if table_name[0].isdigit():
        table_name = (
            "Table_" +
            table_name
        )

    reference = (
        f"A1:"
        f"{get_column_letter(worksheet.max_column)}"
        f"{worksheet.max_row}"
    )

    table = Table(
        displayName=table_name,
        ref=reference
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    table.tableStyleInfo = style

    worksheet.add_table(
        table
    )


# ============================================================
# CONFIDENCE / VALIDATION
# ============================================================

def add_confidence_column(
    worksheet,
    confidence_values,
    header="Confidence"
):
    """
    Add extraction confidence values.

    confidence_values can be:
    [99, 95, 87, ...]
    """

    column = (
        worksheet.max_column + 1
    )

    cell = worksheet.cell(
        row=1,
        column=column
    )

    cell.value = header

    for index, confidence in enumerate(
        confidence_values,
        start=2
    ):

        cell = worksheet.cell(
            row=index,
            column=column
        )

        if confidence is None:
            confidence = ""

        cell.value = confidence

        if isinstance(
            confidence,
            (int, float)
        ):

            cell.number_format = (
                '0.00'
            )

            if confidence < 80:

                cell.comment = Comment(
                    "Low extraction confidence. "
                    "Please verify this row.",
                    "DocuConvert"
                )

    apply_header_style(
        worksheet
    )


def add_validation_column(
    worksheet,
    validation_values,
    header="Validation"
):
    """
    Add validation status per row.
    """

    column = (
        worksheet.max_column + 1
    )

    worksheet.cell(
        row=1,
        column=column
    ).value = header

    for index, status in enumerate(
        validation_values,
        start=2
    ):

        worksheet.cell(
            row=index,
            column=column
        ).value = (
            status
            if status
            else "Not checked"
        )

    apply_header_style(
        worksheet
    )


# ============================================================
# SUMMARY SHEET
# ============================================================

def create_summary_sheet(
    workbook,
    filename=None,
    row_count=0,
    column_count=0,
    document_type=None,
    ocr_used=False,
    average_confidence=None
):
    """
    Create a professional Summary sheet.
    """

    worksheet = workbook.create_sheet(
        "Summary",
        0
    )

    worksheet["A1"] = (
        "DocuConvert - Conversion Summary"
    )

    worksheet["A1"].font = TITLE_FONT

    worksheet.merge_cells(
        "A1:B1"
    )

    summary = [
        ("File", filename or ""),
        ("Document Type", document_type or ""),
        ("Rows", row_count),
        ("Columns", column_count),
        ("OCR Used", "Yes" if ocr_used else "No"),
        (
            "Average Confidence",
            average_confidence
            if average_confidence is not None
            else "Not available"
        ),
        (
            "Generated",
            datetime.now().strftime(
                "%d-%m-%Y %H:%M:%S"
            )
        ),
    ]

    for row_index, (
        label,
        value
    ) in enumerate(
        summary,
        start=3
    ):

        worksheet.cell(
            row=row_index,
            column=1
        ).value = label

        worksheet.cell(
            row=row_index,
            column=2
        ).value = value

        worksheet.cell(
            row=row_index,
            column=1
        ).font = SUBTITLE_FONT

        worksheet.cell(
            row=row_index,
            column=1
        ).border = THIN_BORDER

        worksheet.cell(
            row=row_index,
            column=2
        ).border = THIN_BORDER

    worksheet.column_dimensions[
        "A"
    ].width = 25

    worksheet.column_dimensions[
        "B"
    ].width = 40

    if isinstance(
        average_confidence,
        (int, float)
    ):

        worksheet["B8"].number_format = (
            "0.00"
        )

    return worksheet


# ============================================================
# MAIN EXCEL EXPORT
# ============================================================

def export_table_to_excel(
    rows,
    output_path,
    headers=None,
    sheet_name="Data",
    title=None,
    filename=None,
    document_type=None,
    ocr_used=False,
    confidence_values=None,
    validation_values=None
):
    """
    Main universal Excel export function.

    Parameters
    ----------
    rows:
        List of rows.

    output_path:
        Destination .xlsx path.

    headers:
        Optional column headers.

    sheet_name:
        Excel worksheet name.

    title:
        Optional title above the data.

    filename:
        Original document filename.

    document_type:
        Type of source document.

    ocr_used:
        Whether OCR was used.

    confidence_values:
        Optional row confidence values.

    validation_values:
        Optional row validation statuses.
    """

    if not output_path:
        raise ValueError(
            "Output path is required."
        )

    output_directory = os.path.dirname(
        os.path.abspath(
            output_path
        )
    )

    os.makedirs(
        output_directory,
        exist_ok=True
    )

    # Prepare data
    headers, data_rows = prepare_table(
        rows,
        headers
    )

    # Create workbook
    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = safe_sheet_name(
        sheet_name
    )

    current_row = 1

    # Optional title
    if title:

        worksheet.cell(
            row=current_row,
            column=1
        ).value = title

        worksheet.cell(
            row=current_row,
            column=1
        ).font = TITLE_FONT

        worksheet.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=max(
                len(headers),
                1
            )
        )

        current_row += 2

    # Headers
    for column_index, header in enumerate(
        headers,
        start=1
    ):

        worksheet.cell(
            row=current_row,
            column=column_index
        ).value = header

    header_row = current_row

    current_row += 1

    # Data
    for row in data_rows:

        for column_index, value in enumerate(
            row,
            start=1
        ):

            cell = worksheet.cell(
                row=current_row,
                column=column_index
            )

            cell.value = clean_value(
                value
            )

        current_row += 1

    # Styling
    apply_header_style(
        worksheet,
        header_row
    )

    apply_data_style(
        worksheet,
        header_row + 1
    )

    # Convert data types
    format_data_types(
        worksheet,
        header_row + 1
    )

    # Freeze header
    worksheet.freeze_panes = (
        f"A{header_row + 1}"
    )

    # Filters/table
    if (
        worksheet.max_row >
        header_row
    ):

        # Table starts at actual header row.
        reference = (
            f"A{header_row}:"
            f"{get_column_letter(worksheet.max_column)}"
            f"{worksheet.max_row}"
        )

        table = Table(
            displayName="ConvertedData",
            ref=reference
        )

        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        table.tableStyleInfo = style

        worksheet.add_table(
            table
        )

    # Confidence column
    if confidence_values:

        # Ensure values correspond to data rows
        usable_values = confidence_values[
            :len(data_rows)
        ]

        confidence_column = (
            worksheet.max_column + 1
        )

        worksheet.cell(
            row=header_row,
            column=confidence_column
        ).value = "Confidence"

        for index, confidence in enumerate(
            usable_values,
            start=header_row + 1
        ):

            cell = worksheet.cell(
                row=index,
                column=confidence_column
            )

            cell.value = confidence

            if isinstance(
                confidence,
                (int, float)
            ):

                cell.number_format = (
                    "0.00"
                )

                if confidence < 80:

                    cell.comment = Comment(
                        "Low extraction confidence. "
                        "Please verify this data.",
                        "DocuConvert"
                    )

        worksheet.cell(
            row=header_row,
            column=confidence_column
        ).font = HEADER_FONT

        worksheet.cell(
            row=header_row,
            column=confidence_column
        ).alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Validation column
    if validation_values:

        usable_values = validation_values[
            :len(data_rows)
        ]

        validation_column = (
            worksheet.max_column + 1
        )

        worksheet.cell(
            row=header_row,
            column=validation_column
        ).value = "Validation"

        for index, status in enumerate(
            usable_values,
            start=header_row + 1
        ):

            worksheet.cell(
                row=index,
                column=validation_column
            ).value = (
                status
                if status
                else "Not checked"
            )

        worksheet.cell(
            row=header_row,
            column=validation_column
        ).font = HEADER_FONT

        worksheet.cell(
            row=header_row,
            column=validation_column
        ).alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Conditional formatting for low confidence
    confidence_column = None

    for column in range(
        1,
        worksheet.max_column + 1
    ):

        if worksheet.cell(
            row=header_row,
            column=column
        ).value == "Confidence":

            confidence_column = column
            break

    if confidence_column:

        letter = get_column_letter(
            confidence_column
        )

        worksheet.conditional_formatting.add(
            f"{letter}{header_row + 1}:"
            f"{letter}{worksheet.max_row}",
            CellIsRule(
                operator="lessThan",
                formula=["80"]
            )
        )

    # Formatting
    auto_size_columns(
        worksheet
    )

    apply_row_heights(
        worksheet,
        header_row + 1
    )

    worksheet.row_dimensions[
        header_row
    ].height = 30

    # Summary
    average_confidence = None

    if confidence_values:

        numeric_confidence = [
            value
            for value in confidence_values
            if isinstance(
                value,
                (int, float)
            )
        ]

        if numeric_confidence:

            average_confidence = (
                sum(numeric_confidence)
                /
                len(numeric_confidence)
            )

    create_summary_sheet(
        workbook=workbook,
        filename=filename,
        row_count=len(data_rows),
        column_count=len(headers),
        document_type=document_type,
        ocr_used=ocr_used,
        average_confidence=average_confidence
    )

    # Save
    workbook.save(
        output_path
    )

    return output_path


# ============================================================
# BANK STATEMENT EXPORT
# ============================================================

def export_bank_statement_to_excel(
    transactions,
    output_path,
    filename=None,
    document_type="Bank Statement",
    ocr_used=False
):
    """
    Export structured bank transactions.

    Expected transaction format:

    {
        "date": "...",
        "value_date": "...",
        "branch": "...",
        "reference": "...",
        "description": "...",
        "withdrawal": 5000,
        "deposit": 90000,
        "balance": 169384.94,
        "confidence": 97,
        "validation": "Valid"
    }
    """

    headers = [
        "Date",
        "Value Date",
        "Branch",
        "Ref/Chq No",
        "Description",
        "Withdrawals",
        "Deposits",
        "Balance"
    ]

    rows = []

    confidence_values = []

    validation_values = []

    for transaction in transactions:

        rows.append([
            transaction.get(
                "date",
                ""
            ),
            transaction.get(
                "value_date",
                ""
            ),
            transaction.get(
                "branch",
                ""
            ),
            transaction.get(
                "reference",
                transaction.get(
                    "ref_chq_no",
                    ""
                )
            ),
            transaction.get(
                "description",
                transaction.get(
                    "particulars",
                    ""
                )
            ),
            transaction.get(
                "withdrawal",
                transaction.get(
                    "withdrawals"
                )
            ),
            transaction.get(
                "deposit",
                transaction.get(
                    "deposits"
                )
            ),
            transaction.get(
                "balance"
            )
        ])

        confidence_values.append(
            transaction.get(
                "confidence"
            )
        )

        validation_values.append(
            transaction.get(
                "validation",
                "Not checked"
            )
        )

    return export_table_to_excel(
        rows=rows,
        output_path=output_path,
        headers=headers,
        sheet_name="Transactions",
        title="Bank Statement Transactions",
        filename=filename,
        document_type=document_type,
        ocr_used=ocr_used,
        confidence_values=confidence_values,
        validation_values=validation_values
    )


# ============================================================
# GENERIC DATA EXPORT
# ============================================================

def export_data_to_excel(
    data,
    output_path,
    headers=None,
    sheet_name="Data",
    filename=None,
    document_type=None,
    ocr_used=False
):
    """
    Generic wrapper for exporting structured data.
    """

    return export_table_to_excel(
        rows=data,
        output_path=output_path,
        headers=headers,
        sheet_name=sheet_name,
        filename=filename,
        document_type=document_type,
        ocr_used=ocr_used
    )